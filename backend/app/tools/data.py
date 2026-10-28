import csv
import hashlib
import os
import random
import tempfile
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def prepare_chart(path: Path, category: str, value: str, sheet_name: str | None, chart_sheet: str):
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("Se requiere un Excel .xlsx de hasta 20 MB")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    values = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = values[sheet_name or values.sheetnames[0]]
        if sheet.max_row > 10001 or sheet.max_column > 100:
            raise ValueError("Excel supera límites de lectura")
        rows = iter(sheet.values)
        headers = list(next(rows, []))
        if headers.count(category) != 1 or headers.count(value) != 1:
            raise ValueError("Las columnas deben existir y tener nombres únicos: " + str(headers))
        totals, count = {}, 0
        for row in rows:
            if all(item is None for item in row):
                continue
            label = row[headers.index(category)]
            raw = row[headers.index(value)]
            try:
                amount = Decimal(str(raw))
                if not amount.is_finite() or label is None:
                    raise InvalidOperation()
            except (InvalidOperation, ValueError):
                raise ValueError(f"Fila {count + 2}: categoría vacía o valor no numérico; revise fórmulas sin valor guardado") from None
            label = str(label)
            totals[label] = totals.get(label, Decimal(0)) + amount
            count += 1
        if not totals or len(totals) > 100:
            raise ValueError("El gráfico requiere datos y un máximo de 100 categorías")
        source_sheet = sheet.title
    finally:
        values.close()
    book = load_workbook(path)
    try:
        if chart_sheet in book.sheetnames:
            raise ValueError("La pestaña del gráfico ya existe; indique otro nombre")
        target = book.create_sheet(chart_sheet)
        target.append([category, value])
        for label, amount in sorted(totals.items()):
            target.append([label, float(amount)])
            target.cell(target.max_row, 1).data_type = "s"
        chart = BarChart()
        chart.title = f"{value} por {category}"
        chart.x_axis.title, chart.y_axis.title = category, value
        chart.add_data(Reference(target, min_col=2, min_row=1, max_row=target.max_row), titles_from_data=True)
        chart.set_categories(Reference(target, min_col=1, min_row=2, max_row=target.max_row))
        target.add_chart(chart, "D2")
        target.column_dimensions["A"].width = 30
        target.column_dimensions["B"].width = 20
        return book, digest, {"path": str(path), "source_sheet": source_sheet, "chart_sheet": chart_sheet,
                              "chart_type": "bar", "source_rows": count, "categories": len(totals),
                              "total": float(sum(totals.values())), "category_column": category, "value_column": value}
    except Exception:
        book.close()
        raise


def save_chart(book, path: Path, digest: str):
    # Prepare and validate a separate workbook before replacing the authorized original.
    descriptor, temporary = tempfile.mkstemp(suffix=".xlsx", prefix=".office-chart-", dir=path.parent)
    os.close(descriptor)
    temp = Path(temporary)
    try:
        book.save(temp)
        check = load_workbook(temp)
        try:
            if not check[book.sheetnames[-1]]._charts:
                raise ValueError("No se pudo validar el gráfico")
        finally:
            check.close()
        if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
            raise ValueError("El archivo cambió durante la operación; vuelva a revisar y aprobar")
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def read_csv(path: Path, delimiter: str = ",") -> dict:
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("CSV supera 20 MB")
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter=delimiter)
        if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise ValueError("Cabeceras vacías o duplicadas")
        rows = []
        for row in reader:
            if len(rows) >= 10000:
                raise ValueError("CSV supera 10000 filas; divida el archivo")
            if None in row:
                raise ValueError("CSV con filas de longitud distinta a sus cabeceras")
            rows.append(row)
    return {"rows": rows, "count": len(rows)}


def summarize(rows: list[dict]) -> dict:
    columns = list(dict.fromkeys(key for row in rows for key in row))
    summary = []
    for column in columns:
        values = [row.get(column) for row in rows if row.get(column) not in (None, "")]
        numbers = []
        for value in values:
            try:
                number = Decimal(str(value))
                if not number.is_finite():
                    raise InvalidOperation()
                numbers.append(number)
            except (InvalidOperation, ValueError):
                numbers = []
                break
        item = {
            "columna": column,
            "cantidad": len(values),
            "faltantes": len(rows) - len(values),
            "tipo": "numérica" if numbers else "texto",
        }
        if numbers:
            item.update(
                suma=float(sum(numbers)),
                promedio=float(sum(numbers) / len(numbers)),
                minimo=float(min(numbers)),
                maximo=float(max(numbers)),
            )
        summary.append(item)
    return {"rows": summary, "source_count": len(rows)}


def write_excel(path: Path, rows: list[dict], summary: list[dict] | None = None):
    if path.suffix.lower() != ".xlsx":
        raise ValueError("El archivo debe tener extensión .xlsx")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, data in [("Datos", rows), ("Resumen", summary or summarize(rows)["rows"])]:
        sheet = workbook.create_sheet(title)
        columns = list(dict.fromkeys(key for row in data for key in row))
        if not columns:
            sheet.append(["Sin datos"])
            continue
        sheet.append(columns)
        for row in data:
            values = []
            for column in columns:
                value = row.get(column)
                # Untrusted CSV/model text must never become a spreadsheet formula.
                if isinstance(value, (dict, list)):
                    value = str(value)
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                    value = "'" + value
                values.append(value)
            sheet.append(values)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="25465F")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                45, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
            )
        if data:
            table = Table(displayName=f"Tabla{title}", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        if title == "Resumen" and "suma" in columns and data:
            chart = BarChart()
            chart.title = "Suma por columna numérica"
            idx = columns.index("suma") + 1
            chart.add_data(Reference(sheet, min_col=idx, min_row=1, max_row=sheet.max_row), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row))
            sheet.add_chart(chart, "K2")
    # Exclusive creation prevents a race from accidentally overwriting a file.
    with path.open("xb") as output:
        workbook.save(output)
    check = load_workbook(path, read_only=True, data_only=False)
    try:
        if "Datos" not in check.sheetnames or check["Datos"].max_row != max(1, len(rows) + 1):
            raise ValueError("Validación de Excel falló")
    finally:
        check.close()
    return {"path": str(path), "rows": len(rows), "validated": True}


def generate_sales(count: int, seed: int | None = None) -> list[dict]:
    """Synthetic data generated deterministically with an optional seed, never by the model."""
    if isinstance(count, bool) or not isinstance(count, int) or not 1 <= count <= 10000:
        raise ValueError("La cantidad debe estar entre 1 y 10000")
    rng = random.Random(seed)
    products = [("Cuaderno", 1500), ("Lápiz", 500), ("Carpeta", 2000), ("Agenda", 4500)]
    rows = []
    for index in range(count):
        product, price = rng.choice(products)
        quantity = rng.randint(1, 12)
        rows.append(
            {
                "id": index + 1,
                "fecha": (date(2026, 1, 1) + timedelta(days=rng.randrange(365))).isoformat(),
                "lugar_compra": rng.choice(["Santiago", "Rancagua", "Talca"]),
                "producto": product,
                "cantidad": quantity,
                "precio_unitario": price,
                "total": quantity * price,
                "datos_ficticios": True,
            }
        )
    return rows
