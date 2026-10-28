import json
import time
from pathlib import PureWindowsPath

from jsonschema import ValidationError, validate
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.audit import audit, redact
from app.execution.service import execute as execute_backend
from app.models import Agent, AgentRelation, AgentTool, Artifact, Execution, Task, Tool
from app.security.permissions import ApprovalRequired, require_approval, resource_path
from app.tools.connections import describe_connection, execute_write, query_readonly
from app.tools.data import generate_sales, prepare_chart, read_csv, save_chart, summarize, write_excel
from app.tools.files import publish_bytes


class ToolArgumentsError(ValueError):
    """Rejected arguments before any tool side effect; safe for bounded correction."""


def assigned_tools(session: Session, agent: Agent) -> list[Tool]:
    return list(
        session.scalars(
            select(Tool)
            .join(AgentTool, Tool.id == AgentTool.tool_id)
            .where(AgentTool.agent_id == agent.id, Tool.active.is_(True))
        )
    )


def dispatch(session: Session, task: Task, agent: Agent, tool: Tool, args: dict) -> dict:
    kind = tool.kind
    excel_kinds = {"excel.create", "excel.generate_sales", "excel.from_csv", "sql.export_excel"}
    if kind == "agents.list":
        allowed = set(session.scalars(select(AgentRelation.target_id).where(AgentRelation.source_id == agent.id)))
        return {"agents": [{"id": item.id, "name": item.name, "goal": item.goal, "status": item.status,
                            "delegatable": item.id in allowed and item.status != "disabled"}
                           for item in session.scalars(select(Agent).where(Agent.office_id == agent.office_id))]}
    if kind == "tasks.status":
        target = session.get(Task, args.get("task_id", task.id))
        if not target or target.office_id != agent.office_id:
            raise PermissionError("La tarea no pertenece a esta oficina")
        return {"task_id": target.id, "number": target.number, "status": target.status, "error": target.error,
                "executions": [{"tool": e.tool_name, "status": e.status, "error": e.error}
                               for e in session.scalars(select(Execution).where(Execution.task_id == target.id)
                               .order_by(Execution.created_at.desc()).limit(20))]}
    if kind.startswith("web.") or kind == "api.request":
        from app.tools.network import api_request, web_tool
        return api_request(session, task, agent, args) if kind == "api.request" else web_tool(session, task, agent, kind, args)
    if kind.startswith("pdf."):
        from app.tools.pdf import create_pdf, merge_pdf, read_pdf
        if kind == "pdf.read":
            return read_pdf(session, agent, args)
        return (create_pdf if kind == "pdf.create" else merge_pdf)(session, task, agent, args)
    if kind == "python.packages":
        return execute_backend(session, task, agent, kind, args, tool.timeout)
    if kind == "sports.analyze":
        return execute_backend(session, task, agent, kind, args, tool.timeout)
    if kind == "sql.execute":
        return execute_write(session, task, agent, args)
    if kind == "filesystem.write":
        return publish_bytes(session, task, agent, args["resource_id"], args["path"], args["content"].encode("utf-8"), kind)
    if kind == "excel.chart":
        path = resource_path(session, agent, args["resource_id"], args["path"], "read")
        resource_path(session, agent, args["resource_id"], args["path"], "modify")
        book, digest, result = prepare_chart(path, args["category_column"], args["value_column"],
                                             args.get("sheet"), args.get("chart_sheet", "GraficoVentas"))
        try:
            if agent.mode == "safe":
                require_approval(session, task, agent, "overwrite", {"tool": tool.id, "args": args, "sha256": digest},
                                 f"Agregar pestaña {result['chart_sheet']} y gráfico al archivo {path.name}")
            # Recheck authorized path immediately before replacement.
            resource_path(session, agent, args["resource_id"], args["path"], "modify")
            save_chart(book, path, digest)
        finally:
            book.close()
        session.add(Artifact(task_id=task.id, agent_id=agent.id, resource_id=args["resource_id"],
                             path=str(path), name=path.name))
        return {**result, "validated": True}
    if kind == "social.prepare_post":
        relative_path = args["path"]
        if PureWindowsPath(relative_path).suffix.lower() != ".json":
            raise ValueError("El borrador de publicación debe tener extensión .json")
        path = resource_path(session, agent, args["resource_id"], relative_path, "create")
        draft = {
            "platform": args["platform"],
            "text": args["text"],
            "status": "draft",
            "published": False,
            "requires_publish_approval": True,
        }
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("x", encoding="utf-8") as output:
            json.dump(draft, output, ensure_ascii=False, indent=2)
        session.add(
            Artifact(
                task_id=task.id, agent_id=agent.id, resource_id=args["resource_id"], path=str(path), name=path.name
            )
        )
        return {
            "path": str(path),
            "status": "draft",
            "published": False,
            "next_step": "Revisar contenido; falta conectar y autorizar la publicación mediante API oficial",
        }
    if kind == "filesystem.list":
        folder = resource_path(session, agent, args["resource_id"], args.get("path", ""), "read")
        if not folder.is_dir():
            raise ValueError("El recurso debe ser una carpeta")
        limit = args.get("limit", 100)
        entries = []
        for child in folder.iterdir():
            if child.is_symlink() or (hasattr(child, "is_junction") and child.is_junction()):
                continue
            entries.append({"name": child.name, "kind": "directory" if child.is_dir() else "file"})
            if len(entries) > limit:
                break
        return {"entries": entries[:limit], "truncated": len(entries) > limit}
    if kind in {"filesystem.read", "csv.read", "excel.read"}:
        path = resource_path(session, agent, args["resource_id"], args["path"], "read")
        if kind == "csv.read":
            return read_csv(path, args.get("delimiter", ","))
        if kind == "excel.read":
            from openpyxl import load_workbook

            if path.stat().st_size > 20 * 1024 * 1024:
                raise ValueError("Excel supera 20 MB")
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                result = {}
                for sheet in book:
                    if sheet.max_row > 10001 or sheet.max_column > 100:
                        raise ValueError("Excel supera límites de lectura")
                    result[sheet.title] = [
                        [str(v) if not isinstance(v, (str, float, int, bool, type(None))) else v for v in row]
                        for row in sheet.values
                    ]
                return {"sheets": result}
            finally:
                book.close()
        if path.stat().st_size > 2 * 1024 * 1024:
            raise ValueError("Texto supera 2 MB")
        return {"content": path.read_text(encoding="utf-8")}
    if kind in {"filesystem.write", "filesystem.create_directory"} | excel_kinds:
        relative_path = args["path"]
        if kind in excel_kinds:
            suffix = PureWindowsPath(relative_path).suffix
            if not suffix:
                relative_path += ".xlsx"
            elif suffix.lower() == ".xlxs":
                relative_path = relative_path[: -len(suffix)] + ".xlsx"
            elif suffix.lower() != ".xlsx":
                raise ValueError("El archivo debe tener extensión .xlsx")
        path = resource_path(session, agent, args["resource_id"], relative_path, "create")
        if kind == "filesystem.create_directory":
            path.mkdir(parents=True, exist_ok=True)
            return {"path": str(path)}
        if path.exists():
            resource_path(session, agent, args["resource_id"], relative_path, "modify")
            if agent.mode == "safe":
                require_approval(
                    session,
                    task,
                    agent,
                    "overwrite",
                    {"tool": tool.id, "args": args},
                    f"Sobrescribir archivo {path.name}",
                )
            # Excel writer uses exclusive creation; intentionally refuse overwrites for now.
            if kind in excel_kinds:
                raise ValueError("El Excel ya existe. Use un nombre nuevo")
        path.parent.mkdir(parents=True, exist_ok=True)
        if kind in excel_kinds:
            if kind == "excel.generate_sales":
                rows = generate_sales(args["count"], args.get("seed"))
            elif kind == "excel.from_csv":
                source = resource_path(session, agent, args["source_resource_id"], args["source_path"], "read")
                rows = read_csv(source, args.get("delimiter", ","))["rows"]
            elif kind == "sql.export_excel":
                rows = query_readonly(session, agent, args["connection_id"], args["query"])["rows"]
            else:
                rows = args["rows"]
            result = write_excel(path, rows, args.get("summary"))
            if kind == "excel.generate_sales":
                result["synthetic"] = True
        else:
            with path.open("w" if path.exists() else "x", encoding="utf-8") as output:
                output.write(args["content"])
            result = {"path": str(path)}
        session.add(
            Artifact(
                task_id=task.id, agent_id=agent.id, resource_id=args["resource_id"], path=str(path), name=path.name
            )
        )
        return result
    if kind == "data.summarize":
        return summarize(args["rows"])
    if kind == "python.execute":
        return execute_backend(session, task, agent, kind, args, tool.timeout)
    if kind == "sql.describe":
        return describe_connection(session, agent, args["connection_id"], args.get("table"), args.get("schema"))
    if kind == "sql.query":
        return query_readonly(session, agent, args["connection_id"], args["query"], args.get("parameters"), args.get("max_rows", 10000))
    if kind == "custom.sql":
        return query_readonly(session, agent, tool.config["connection_id"], tool.config["query"])
    if kind == "custom.python":
        return execute_backend(session, task, agent, "python.execute",
                               {"code": tool.config["code"], "resources": tool.config.get("resources", [])}, tool.timeout)
    raise ValueError("Tipo de herramienta pendiente: no se ejecutará")


def execute_tool(session: Session, task: Task, agent: Agent, name: str, args: dict):
    tool = next((t for t in assigned_tools(session, agent) if t.name == name), None)
    if not tool:
        raise PermissionError("Herramienta no asignada o desactivada")
    if not set(tool.required_permissions).issubset(agent.permissions):
        raise PermissionError("Permisos de herramienta no concedidos al agente")
    execution = Execution(task_id=task.id, agent_id=agent.id, tool_name=name, input=redact(args))
    session.add(execution)
    session.commit()
    start = time.monotonic()
    audit(session, 'tool.started', {'name': name}, task.id, agent.id)
    error_category = None
    try:
        if tool.kind == 'web.fetch' and isinstance(args.get('destination'), dict):
            path = args['destination'].get('path', '')
            if isinstance(path, str) and not path.lower().endswith('.txt'):
                raise ToolArgumentsError(
                    'web.fetch extrae texto plano de una página y destination.path debe terminar en .txt. '
                    'No guarda HTML original: cambia .html/.htm por .txt si necesitas leer la página. '
                    'Para un dataset CSV/JSON/XLSX usa web.download con filename y destination como carpeta; '
                    'para las API configuradas usa api.request con destination.path terminado en .json.'
                )
        if tool.kind == "excel.create" and "format" in args:
            raise ToolArgumentsError(
                "excel.create crea archivos .xlsx y no acepta el parámetro format. "
                "Use un nombre relativo como ventas26.xlsx y filas con datos. "
                "Para CSV use filesystem.write si está asignada, con path y content."
            )
        try:
            validate(args, tool.parameters)
        except ValidationError as exc:
            location = ".".join(str(part) for part in exc.absolute_path) or "arguments"
            raise ToolArgumentsError(
                f"Argumentos inválidos para {name} en {location} ({exc.validator}); "
                "use únicamente los campos y tipos del esquema de la herramienta"
            ) from exc
        result = dispatch(session, task, agent, tool, args)
        execution.status, execution.output = "completed", redact(result)
        return result
    except ApprovalRequired:
        execution.status = "waiting_approval"
        error_category = "approval_required"
        raise
    except Exception as exc:
        execution.status, execution.error = "error", redact(str(exc))
        error_category = ("permission" if isinstance(exc, PermissionError) else
                          "tool_timeout" if isinstance(exc, TimeoutError) else
                          "tool_arguments" if isinstance(exc, ToolArgumentsError) else "tool_error")
        raise
    finally:
        execution.duration_ms = int((time.monotonic() - start) * 1000)
        session.commit()
        audit(
            session,
            "tool." + execution.status,
            {"name": name, "input": args, "output": execution.output, "error": execution.error,
             "duration_ms": execution.duration_ms,
             "error_category": error_category},
            task.id,
            agent.id,
        )
