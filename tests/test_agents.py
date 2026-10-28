"""Contract fixtures are test-only. No production provider is replaced."""

import asyncio

import pytest
from app.ai.providers import AIProvider, AIResponse, parse_response, providers, validate_provider
from app.core.db import SessionLocal
from app.models import Memory, Provider
from conftest import create, wait_task
from sqlalchemy import select


class ScriptedProvider(AIProvider):
    def __init__(self, actions):
        self.actions = actions
        self.calls = []

    async def complete(self, provider, agent, model, messages):
        self.calls.append((provider.kind, agent.id, model, messages))
        action = self.actions[agent.id].pop(0)
        if isinstance(action, Exception):
            raise action
        return AIResponse(action, 20, 10)


def test_local_delegation_and_isolated_memory(client, monkeypatch):
    office = create(client, "offices", name="Dinámica")
    provider = client.get("/api/providers").json()[0]
    parent = create(
        client,
        "agents",
        name="Coordinador libre",
        office_id=office["id"],
        provider_id=provider["id"],
        model="test-only-model",
        memory_mode="persistent",
    )
    child = create(
        client,
        "agents",
        name="Especialista libre",
        office_id=office["id"],
        provider_id=provider["id"],
        model="test-only-model",
        memory_mode="persistent",
    )
    create(client, "relations", source_id=parent["id"], target_id=child["id"])
    fixture = ScriptedProvider(
        {
            parent["id"]: [
                {"type": "delegate", "agent_id": child["id"], "request": "Calcula"},
                {"type": "final", "content": "Resultado coordinado"},
            ],
            child["id"]: [{"type": "final", "content": "Resultado privado del especialista"}],
        }
    )
    monkeypatch.setitem(providers, "ollama", fixture)
    task = create(client, "tasks", office_id=office["id"], agent_id=parent["id"], request="Coordina")
    assert wait_task(client, task["id"])["status"] == "completed"
    details = client.get(f"/api/tasks/{task['id']}/details").json()
    assert len(details["messages"]) == 2
    scene = client.get(f"/api/offices/{office['id']}/scene").json()
    assert {message["id"] for message in scene["messages"]} == {message["id"] for message in details["messages"]}
    assert all("content" not in message for message in scene["messages"])
    with SessionLocal() as session:
        memories = list(session.scalars(select(Memory)))
        assert {m.agent_id for m in memories} == {parent["id"], child["id"]}
    assert "Resultado privado" not in fixture.calls[0][3][0]["content"]


def test_external_ask_blocks_before_network_and_resumes(client, monkeypatch):
    office = create(client, "offices", name="O")
    provider = create(
        client,
        "providers",
        name="OpenAI configurado",
        kind="openai",
        url="https://api.openai.com/v1",
        default_model="test-model",
    )
    agent = create(
        client,
        "agents",
        name="A",
        office_id=office["id"],
        provider_id=provider["id"],
        external_policy="ask",
        model="test-model",
    )
    fixture = ScriptedProvider({agent["id"]: [{"type": "final", "content": "Externo autorizado"}]})
    monkeypatch.setitem(providers, "openai", fixture)
    task = create(client, "tasks", office_id=office["id"], agent_id=agent["id"], request="Consulta")
    assert wait_task(client, task["id"])["status"] == "waiting_approval"
    assert len(fixture.calls) == 0
    approval = client.get(f"/api/tasks/{task['id']}/details").json()["approvals"][0]
    assert "costos externos" in approval["reason"]
    assert client.post(f"/api/approvals/{approval['id']}/decide", json={"approved": True}).status_code == 200
    assert wait_task(client, task["id"])["status"] == "completed"
    assert len(fixture.calls) == 1


def test_external_never_blocks_fallback(client, monkeypatch):
    office = create(client, "offices", name="O")
    local = client.get("/api/providers").json()[0]
    external = create(client, "providers", name="External", kind="openai", url="https://api.openai.com/v1")
    agent = create(
        client,
        "agents",
        name="A",
        office_id=office["id"],
        provider_id=local["id"],
        model="test-model",
        external_policy="never",
        fallbacks=[{"provider_id": external["id"], "model": "test-model"}],
    )
    local_fixture = ScriptedProvider({agent["id"]: [ValueError("Ollama desconectado")]})
    external_fixture = ScriptedProvider({agent["id"]: []})
    monkeypatch.setitem(providers, "ollama", local_fixture)
    monkeypatch.setitem(providers, "openai", external_fixture)
    task = create(client, "tasks", office_id=office["id"], agent_id=agent["id"], request="Haz trabajo")
    finished = wait_task(client, task["id"])
    assert finished["status"] == "error"
    assert "prohibidas" in finished["error"]
    assert external_fixture.calls == []


def test_agent_tool_checks_and_real_result(client, tmp_path, monkeypatch):
    office = create(client, "offices", name="O")
    provider = client.get("/api/providers").json()[0]
    agent = create(client, "agents", name="A", office_id=office["id"], provider_id=provider["id"], model="test-model")
    resource = create(client, "resources", name="R", path=str(tmp_path))
    tool = next(t for t in client.get("/api/tools").json() if t["name"] == "filesystem.write")
    assert (
        client.put(
            f"/api/agents/{agent['id']}/assignments",
            json={
                "tools": [tool["id"]],
                "resources": [{"resource_id": resource["id"], "permissions": ["read", "create"]}],
            },
        ).status_code
        == 200
    )
    fixture = ScriptedProvider(
        {
            agent["id"]: [
                {
                    "type": "tool",
                    "name": "filesystem.write",
                    "arguments": {"resource_id": resource["id"], "path": "actual.txt", "content": "real"},
                },
                {"type": "final", "content": "Archivo creado"},
            ]
        }
    )
    monkeypatch.setitem(providers, "ollama", fixture)
    task = create(client, "tasks", office_id=office["id"], agent_id=agent["id"], request="Crea texto")
    assert wait_task(client, task["id"])["status"] == "completed"
    assert (tmp_path / "actual.txt").read_text() == "real"


def test_cancellation_keeps_worker_available(client, monkeypatch):
    office = create(client, "offices", name="O")
    provider = client.get("/api/providers").json()[0]
    agent = create(client, "agents", name="A", office_id=office["id"], provider_id=provider["id"], model="test-model")

    class SlowProvider(AIProvider):
        async def complete(self, provider, agent, model, messages):
            await asyncio.sleep(0.2)
            return AIResponse({"type": "final", "content": "Hecho"})

    monkeypatch.setitem(providers, "ollama", SlowProvider())
    first = create(client, "tasks", office_id=office["id"], agent_id=agent["id"], request="Primera")
    assert client.post(f"/api/tasks/{first['id']}/cancel").status_code == 200
    second = create(client, "tasks", office_id=office["id"], agent_id=agent["id"], request="Segunda")
    assert wait_task(client, second["id"])["status"] == "completed"
    assert client.get(f"/api/tasks/{first['id']}").json()["status"] == "cancelled"


def test_provider_protocol_and_cloud_policy():
    import pytest

    assert parse_response('```json\n{"type":"final","content":"ok"}\n```')["type"] == "final"
    with pytest.raises(ValueError):
        parse_response('{"type":"execute_arbitrary"}')
    with pytest.raises(ValueError):
        validate_provider(Provider(kind="ollama", url="http://127.0.0.1:11434"), "anything:cloud")
    with pytest.raises(ValueError):
        validate_provider(Provider(kind="openai", url="https://other.example/v1"))


@pytest.mark.parametrize("requested", ["ventas26.xlxs", "ventas26", "ventas26.xlsx"])
def test_excel_normalizes_filename_and_records_actual_artifact(client, tmp_path, monkeypatch, requested):
    from openpyxl import load_workbook

    office = create(client, "offices", name="Excel")
    provider = client.get("/api/providers").json()[0]
    agent = create(client, "agents", name="A", office_id=office["id"], provider_id=provider["id"], model="test-model")
    resource = create(client, "resources", name="R", path=str(tmp_path))
    tool = next(t for t in client.get("/api/tools").json() if t["name"] == "excel.create")
    response = client.put(
        f"/api/agents/{agent['id']}/assignments",
        json={
            "tools": [tool["id"]],
            "resources": [{"resource_id": resource["id"], "permissions": ["create"]}],
        },
    )
    assert response.status_code == 200
    fixture = ScriptedProvider(
        {
            agent["id"]: [
                {
                    "type": "tool",
                    "name": "excel.create",
                    "arguments": {"resource_id": resource["id"], "path": requested, "rows": [{"venta": 125}]},
                },
                {"type": "final", "content": "Creado"},
            ]
        }
    )
    monkeypatch.setitem(providers, "ollama", fixture)
    task = create(client, "tasks", office_id=office["id"], agent_id=agent["id"], request="Crea Excel")
    assert wait_task(client, task["id"])["status"] == "completed"
    path = tmp_path / "ventas26.xlsx"
    book = load_workbook(path)
    assert book["Datos"]["A2"].value == 125
    book.close()
    from app.models import Artifact

    with SessionLocal() as session:
        artifact = session.scalar(select(Artifact).where(Artifact.task_id == task["id"]))
        assert artifact.path == str(path)
    # A typo must not bypass modify permission on the normalized destination.
    fixture.actions[agent["id"]] = [
        {
            "type": "tool",
            "name": "excel.create",
            "arguments": {"resource_id": resource["id"], "path": "ventas26.xlxs", "rows": [{"venta": 999}]},
        }
    ]
    retry = create(client, "tasks", office_id=office["id"], agent_id=agent["id"], request="Repite")
    assert wait_task(client, retry["id"])["status"] == "error"
    book = load_workbook(path)
    assert book["Datos"]["A2"].value == 125
    book.close()


def test_excel_format_error_is_readable_and_logged(client, tmp_path, monkeypatch):
    from app.models import Execution

    office = create(client, "offices", name="CSV inválido")
    provider = client.get("/api/providers").json()[0]
    agent = create(client, "agents", name="A", office_id=office["id"], provider_id=provider["id"], model="test-model")
    resource = create(client, "resources", name="R", path=str(tmp_path))
    tool = next(t for t in client.get("/api/tools").json() if t["name"] == "excel.create")
    client.put(
        f"/api/agents/{agent['id']}/assignments",
        json={
            "tools": [tool["id"]],
            "resources": [{"resource_id": resource["id"], "permissions": ["create"]}],
        },
    )
    fixture = ScriptedProvider(
        {
            agent["id"]: [
                {
                    "type": "tool",
                    "name": "excel.create",
                    "arguments": {
                        "resource_id": resource["id"],
                        "path": "/ventas26.csv",
                        "rows": [],
                        "summary": [],
                        "format": "csv",
                    },
                }
            ]
        }
    )
    fixture.actions[agent["id"]] *= 3
    monkeypatch.setitem(providers, "ollama", fixture)
    task = create(client, "tasks", office_id=office["id"], agent_id=agent["id"], request="Crea archivo")
    finished = wait_task(client, task["id"])
    assert finished["status"] == "error"
    assert "no acepta el parámetro format" in finished["error"]
    assert "Failed validating" not in finished["error"]
    with SessionLocal() as session:
        execution = session.scalar(
            select(Execution).where(Execution.task_id == task["id"], Execution.tool_name == "excel.create")
        )
        assert execution.status == "error"
    activity = client.get(f"/api/offices/{office['id']}/scene").json()["agents"][0]["activity"]
    assert activity["tool"] == "excel.create"
    assert "no acepta" in activity["error"]
    assert list(tmp_path.iterdir()) == []


def test_invalid_arguments_corrected_without_repeating_a_write(client, tmp_path, monkeypatch):
    office = create(client, "offices", name="Corrección")
    provider = client.get("/api/providers").json()[0]
    agent = create(
        client,
        "agents",
        name="A",
        office_id=office["id"],
        provider_id=provider["id"],
        model="test-model",
        context_limit=8192,
    )
    resource = create(client, "resources", name="R", path=str(tmp_path))
    tool = next(t for t in client.get("/api/tools").json() if t["name"] == "excel.generate_sales")
    client.put(
        f"/api/agents/{agent['id']}/assignments",
        json={"tools": [tool["id"]], "resources": [{"resource_id": resource["id"], "permissions": ["create"]}]},
    )
    fixture = ScriptedProvider(
        {
            agent["id"]: [
                {
                    "type": "tool",
                    "name": "excel.generate_sales",
                    "arguments": {"resource_id": resource["id"], "path": "ventas.xlsx", "count": 0},
                },
                {
                    "type": "tool",
                    "name": "excel.generate_sales",
                    "arguments": {"resource_id": resource["id"], "path": "ventas.xlsx", "count": 50},
                },
                {"type": "final", "content": "Generadas 50 ventas"},
            ]
        }
    )
    monkeypatch.setitem(providers, "ollama", fixture)
    task = create(client, "tasks", office_id=office["id"], agent_id=agent["id"], request="Genera ventas")
    assert wait_task(client, task["id"])["status"] == "completed"
    assert len(list(tmp_path.iterdir())) == 1
    assert len(fixture.calls) == 3
