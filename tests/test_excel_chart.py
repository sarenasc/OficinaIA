import json

import pytest
from app.agents.engine import model_tool_result
from app.ai.providers import providers
from app.tools.data import generate_sales, prepare_chart, save_chart, write_excel
from conftest import create, wait_task
from openpyxl import load_workbook
from test_agents import ScriptedProvider


def test_chart_preserves_source_and_aggregates_all_rows(tmp_path):
    path = tmp_path / 'ventas.xlsx'
    rows = generate_sales(100, 7)
    write_excel(path, rows)
    book, digest, result = prepare_chart(path, 'producto', 'total', 'Datos', 'GraficoVentas')
    save_chart(book, path, digest)
    book.close()
    check = load_workbook(path)
    assert check['Datos'].max_row == 101
    assert list(check['Datos'].values)[1][0] == rows[0]['id']
    assert check['GraficoVentas']._charts
    assert sum(row[1] for row in list(check['GraficoVentas'].values)[1:]) == sum(r['total'] for r in rows)
    assert result['source_rows'] == 100
    check.close()
    with pytest.raises(ValueError, match='ya existe'):
        prepare_chart(path, 'producto', 'total', 'Datos', 'GraficoVentas')


def test_changed_source_is_never_replaced(tmp_path):
    path = tmp_path / 'ventas.xlsx'
    write_excel(path, generate_sales(10, 1))
    book, digest, _ = prepare_chart(path, 'producto', 'total', 'Datos', 'GraficoVentas')
    changed = path.read_bytes() + b'changed'
    path.write_bytes(changed)
    with pytest.raises(ValueError, match='cambió'):
        save_chart(book, path, digest)
    book.close()
    assert path.read_bytes() == changed
    assert not list(tmp_path.glob('.office-chart-*'))


def test_model_receives_preview_not_all_rows():
    full = {'sheets': {'Datos': [['producto', 'total']] + [['Producto', i] for i in range(1000)]}}
    result = model_tool_result(full)
    assert result['preview_only']
    assert result['sheets']['Datos']['row_count'] == 1000
    assert len(result['sheets']['Datos']['sample']) == 2
    assert len(json.dumps(result)) < 1000
    assert len(full['sheets']['Datos']) == 1001


def test_excel_read_fits_small_agent_context_and_keeps_full_audit(client, monkeypatch, tmp_path):
    office = create(client, 'offices', name='Contexto reducido')
    provider = client.get('/api/providers').json()[0]
    agent = create(client, 'agents', name='Analista', office_id=office['id'],
                   provider_id=provider['id'], model='test-only-model', context_limit=2048)
    resource = create(client, 'resources', name='Ventas', path=str(tmp_path))
    tool = next(t for t in client.get('/api/tools').json() if t['name'] == 'excel.read')
    client.put(f"/api/agents/{agent['id']}/assignments", json={
        'tools': [tool['id']], 'resources': [{'resource_id': resource['id'], 'permissions': ['read']}]})
    write_excel(tmp_path / 'ventas.xlsx', generate_sales(100, 1))
    fixture = ScriptedProvider({agent['id']: [
        {'type': 'tool', 'name': 'excel.read', 'arguments': {'resource_id': resource['id'], 'path': 'ventas.xlsx'}},
        {'type': 'final', 'content': 'Lectura verificada; solo una muestra disponible en contexto.'}]})
    monkeypatch.setitem(providers, 'ollama', fixture)
    task = create(client, 'tasks', office_id=office['id'], agent_id=agent['id'], request='Lee ventas.xlsx')
    assert wait_task(client, task['id'])['status'] == 'completed'
    details = client.get(f"/api/tasks/{task['id']}/details").json()
    execution = next(e for e in details['executions'] if e['tool_name'] == 'excel.read')
    assert len(execution['output']['sheets']['Datos']) == 101
    assert any('preview_only' in m['content'] for m in fixture.calls[-1][3])


def test_chart_requires_approval_before_changing_existing_file(client, tmp_path):
    office = create(client, 'offices', name='Graficos')
    agent = create(client, 'agents', name='Analista', office_id=office['id'])
    resource = create(client, 'resources', name='Ventas', path=str(tmp_path))
    tool = next(t for t in client.get('/api/tools').json() if t['name'] == 'excel.chart')
    client.put(f"/api/agents/{agent['id']}/assignments", json={
        'tools': [tool['id']], 'resources': [{'resource_id': resource['id'], 'permissions': ['read', 'modify']}]})
    path = tmp_path / 'ventas.xlsx'
    write_excel(path, generate_sales(100, 4))
    before = path.read_bytes()
    procedure = create(client, 'procedures', name='Graficar', office_id=office['id'], steps=[{
        'agent_id': agent['id'], 'tool': 'excel.chart', 'arguments': {
            'resource_id': resource['id'], 'path': 'ventas.xlsx', 'category_column': 'producto', 'value_column': 'total'}}])
    task = create(client, 'tasks', office_id=office['id'], agent_id=agent['id'], procedure_id=procedure['id'], request='Graficar')
    assert wait_task(client, task['id'])['status'] == 'waiting_approval'
    assert path.read_bytes() == before
    approval = client.get(f"/api/tasks/{task['id']}/details").json()['approvals'][0]
    assert client.post(f"/api/approvals/{approval['id']}/decide", json={'approved': True}).status_code == 200
    assert wait_task(client, task['id'])['status'] == 'completed'
    check = load_workbook(path)
    assert check['GraficoVentas']._charts
    check.close()
