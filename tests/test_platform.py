import sqlite3

import pytest
from app.core.audit import redact
from app.core.db import SessionLocal
from app.models import Agent, Office
from app.security.permissions import resource_path
from app.tools.connections import validate_select
from conftest import create, wait_task
from openpyxl import load_workbook


def test_auth_and_origin(client):
    assert client.get("/api/offices", headers={"Authorization": "Bearer wrong"}).status_code == 401
    assert (
        client.post("/api/offices", json={"name": "X"}, headers={"Origin": "https://evil.example"}).status_code == 403
    )
    assert client.post("/api/offices", json={"name": ""}).status_code == 422


def test_crud_persistence_and_cycle(client):
    office = create(client, "offices", name="Mi oficina sin rol predefinido")
    a = create(client, "agents", name="Primero", office_id=office["id"])
    b = create(client, "agents", name="Segundo", office_id=office["id"])
    create(client, "relations", source_id=a["id"], target_id=b["id"])
    assert client.post("/api/relations", json={"source_id": b["id"], "target_id": a["id"]}).status_code == 422
    assert client.patch("/api/agents/" + a["id"], json={"position": {"x": 20, "y": 70}}).status_code == 200
    with SessionLocal() as session:
        assert session.get(Office, office["id"]).name == office["name"]
        assert session.get(Agent, a["id"]).position == {"x": 20, "y": 70}
    other = create(client, "offices", name="Otra")
    c = create(client, "agents", name="Otro", office_id=other["id"])
    assert client.post("/api/relations", json={"source_id": a["id"], "target_id": c["id"]}).status_code == 422


def test_demo_excel_end_to_end_and_removal(client, tmp_path):
    demo_response = client.post("/api/demo/install", json={"output_path": str(tmp_path)})
    assert demo_response.status_code == 200, demo_response.text
    demo = demo_response.json()
    task = create(
        client,
        "tasks",
        office_id=demo["office"]["id"],
        agent_id=demo["agent_id"],
        request="Generar resumen",
        procedure_id=demo["procedure"]["id"],
        inputs={"output_name": "real.xlsx"},
    )
    finished = wait_task(client, task["id"])
    assert finished["status"] == "completed", finished
    assert "validated" in finished["result"]
    output = tmp_path / "real.xlsx"
    assert output.exists()
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Datos", "Resumen"]
    assert workbook["Datos"].max_row == 5
    summary = list(workbook["Resumen"].values)
    kilos = next(row for row in summary[1:] if row[0] == "kilos")
    assert kilos[4] == 6600
    workbook.close()
    details = client.get(f"/api/tasks/{task['id']}/details").json()
    assert len(details["executions"]) == 3
    assert len(details["artifacts"]) == 1
    assert client.get(f"/api/artifacts/{details['artifacts'][0]['id']}/download").status_code == 200
    assert client.delete("/api/demo/remove").status_code == 200
    assert client.get("/api/offices").json() == []
    assert client.get("/api/agents").json() == []
    assert output.exists()  # Never delete user output while removing optional demo.


def test_filesystem_permissions_and_escape(client, tmp_path):
    office = create(client, "offices", name="O")
    agent = create(client, "agents", name="A", office_id=office["id"])
    resource = create(client, "resources", name="R", path=str(tmp_path))
    with SessionLocal() as session:
        with pytest.raises(PermissionError):
            resource_path(session, session.get(Agent, agent["id"]), resource["id"], "file.txt", "read")
    response = client.put(
        f"/api/agents/{agent['id']}/assignments",
        json={"resources": [{"resource_id": resource["id"], "permissions": ["read"]}]},
    )
    assert response.status_code == 200
    with SessionLocal() as session:
        current = session.get(Agent, agent["id"])
        for path in ["../outside.txt", "C:\\Windows\\system.ini", "file.txt:stream"]:
            with pytest.raises(PermissionError):
                resource_path(session, current, resource["id"], path, "read")
        with pytest.raises(PermissionError):
            resource_path(session, current, resource["id"], "new.txt", "create")


@pytest.mark.parametrize(
    "query",
    [
        "DELETE FROM x",
        "SELECT * INTO new_table FROM x",
        "SELECT 1; DROP TABLE x",
        "WITH a AS (DELETE FROM x RETURNING *) SELECT * FROM a",
        "SELECT load_extension('bad')",
        "SELECT pg_read_file('/etc/passwd')",
    ],
)
def test_sql_rejects_write_and_side_effects(query):
    with pytest.raises((PermissionError, ValueError)):
        validate_select(query, "postgresql")


def test_sql_readonly_to_excel(client, tmp_path):
    source = tmp_path / "source.db"
    with sqlite3.connect(source) as db:
        db.execute("CREATE TABLE sales (week INTEGER, kg REAL)")
        db.executemany("INSERT INTO sales VALUES (?,?)", [(1, 20), (2, 30)])
    office = create(client, "offices", name="SQL oficina")
    agent = create(client, "agents", name="Agente configurable", office_id=office["id"])
    connection = create(client, "connections", name="SQLite prueba", kind="sqlite", config={"path": str(source)})
    resource = create(client, "resources", name="Salida", path=str(tmp_path))
    tools = client.get("/api/tools").json()
    assert (
        client.put(
            f"/api/agents/{agent['id']}/assignments",
            json={
                "tools": [t["id"] for t in tools if t["name"] in ["sql.query", "excel.create"]],
                "connections": [{"connection_id": connection["id"], "permissions": ["connect", "select"]}],
                "resources": [{"resource_id": resource["id"], "permissions": ["read", "create"]}],
            },
        ).status_code
        == 200
    )
    procedure = create(
        client,
        "procedures",
        name="SQL → Excel",
        office_id=office["id"],
        steps=[
            {
                "agent_id": agent["id"],
                "tool": "sql.query",
                "arguments": {"connection_id": connection["id"], "query": "SELECT week, kg FROM sales"},
            },
            {
                "agent_id": agent["id"],
                "tool": "excel.create",
                "arguments": {"resource_id": resource["id"], "path": "sql.xlsx", "rows": "$steps.0.rows"},
            },
        ],
    )
    task = create(
        client,
        "tasks",
        office_id=office["id"],
        agent_id=agent["id"],
        request="Consulta de lectura",
        procedure_id=procedure["id"],
    )
    finished = wait_task(client, task["id"])
    assert finished["status"] == "completed", finished
    workbook = load_workbook(tmp_path / "sql.xlsx")
    assert workbook["Datos"]["B3"].value == 30
    workbook.close()


def test_overwrite_approval_and_resume_does_not_repeat_steps(client, tmp_path):
    office = create(client, "offices", name="O")
    agent = create(client, "agents", name="A", office_id=office["id"])
    resource = create(client, "resources", name="R", path=str(tmp_path))
    (tmp_path / "existing.txt").write_text("before")
    tools = client.get("/api/tools").json()
    assert (
        client.put(
            f"/api/agents/{agent['id']}/assignments",
            json={
                "tools": [t["id"] for t in tools if t["name"] == "filesystem.write"],
                "resources": [{"resource_id": resource["id"], "permissions": ["read", "create", "modify"]}],
            },
        ).status_code
        == 200
    )
    procedure = create(
        client,
        "procedures",
        name="Escritura",
        office_id=office["id"],
        steps=[
            {
                "agent_id": agent["id"],
                "tool": "filesystem.write",
                "arguments": {"resource_id": resource["id"], "path": "first.txt", "content": "once"},
            },
            {
                "agent_id": agent["id"],
                "tool": "filesystem.write",
                "arguments": {"resource_id": resource["id"], "path": "existing.txt", "content": "after"},
            },
        ],
    )
    task = create(
        client, "tasks", office_id=office["id"], agent_id=agent["id"], request="Escribir", procedure_id=procedure["id"]
    )
    assert wait_task(client, task["id"])["status"] == "waiting_approval"
    assert (tmp_path / "existing.txt").read_text() == "before"
    details = client.get(f"/api/tasks/{task['id']}/details").json()
    approval = details["approvals"][0]
    assert client.post(f"/api/approvals/{approval['id']}/decide", json={"approved": True}).status_code == 200
    assert wait_task(client, task["id"])["status"] == "completed"
    assert (tmp_path / "existing.txt").read_text() == "after"
    details = client.get(f"/api/tasks/{task['id']}/details").json()
    first = [e for e in details["executions"] if e["input"].get("path") == "first.txt"]
    assert len(first) == 1


def test_secrets_never_in_generic_config(client):
    response = client.post("/api/connections", json={"name": "X", "config": {"password": "hidden"}})
    assert response.status_code == 422
    assert "hidden" not in response.text
    assert (
        redact({"Authorization": "Bearer secret", "nested": {"api_key": "hidden"}})["nested"]["api_key"] == "[REDACTED]"
    )


def test_formula_injection_is_text(tmp_path):
    from app.tools.data import write_excel

    write_excel(tmp_path / "safe.xlsx", [{"x": '=WEBSERVICE("https://bad")'}])
    workbook = load_workbook(tmp_path / "safe.xlsx")
    assert workbook["Datos"]["A2"].data_type == "s"
    workbook.close()
