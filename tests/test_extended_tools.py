import json
import sqlite3

import pytest
from app.tools.data import generate_sales
from conftest import create, wait_task
from openpyxl import load_workbook


def setup_agent(client, tmp_path):
    office = create(client, "offices", name="Herramientas ampliadas")
    agent = create(client, "agents", name="Especialista", office_id=office["id"])
    resource = create(client, "resources", name="Datos", path=str(tmp_path))
    tools = client.get("/api/tools").json()
    client.put(
        f"/api/agents/{agent['id']}/assignments",
        json={
            "tools": [
                t["id"]
                for t in tools
                if t["name"]
                in {
                    "excel.generate_sales",
                    "filesystem.list",
                    "excel.from_csv",
                    "sql.export_excel",
                    "sql.describe",
                    "social.prepare_post",
                }
            ],
            "resources": [{"resource_id": resource["id"], "permissions": ["read", "create"]}],
        },
    )
    return office, agent, resource


def run_tool(client, office, agent, name, arguments):
    procedure = create(
        client,
        "procedures",
        name=name,
        office_id=office["id"],
        steps=[{"agent_id": agent["id"], "tool": name, "arguments": arguments}],
    )
    task = create(
        client,
        "tasks",
        office_id=office["id"],
        agent_id=agent["id"],
        procedure_id=procedure["id"],
        request="Verificación de herramienta",
    )
    return wait_task(client, task["id"])


def test_generated_sales_have_correct_totals_and_repeatable_seed():
    rows = generate_sales(1000, 12)
    assert rows == generate_sales(1000, 12)
    assert len({r["id"] for r in rows}) == 1000
    assert all(r["total"] == r["cantidad"] * r["precio_unitario"] for r in rows)
    assert all(r["datos_ficticios"] for r in rows)
    for count in [0, 10001, True]:
        with pytest.raises(ValueError):
            generate_sales(count)


def test_generator_creates_real_workbook_without_model_rows(client, tmp_path):
    office, agent, resource = setup_agent(client, tmp_path)
    args = {"resource_id": resource["id"], "path": "ventas.xlsx", "count": 50, "seed": 7}
    result = run_tool(client, office, agent, "excel.generate_sales", args)
    assert result["status"] == "completed", result
    book = load_workbook(tmp_path / "ventas.xlsx")
    assert book["Datos"].max_row == 51
    book.close()
    before = (tmp_path / "ventas.xlsx").read_bytes()
    result = run_tool(client, office, agent, "excel.generate_sales", args)
    assert result["status"] == "error"
    assert (tmp_path / "ventas.xlsx").read_bytes() == before
    args["path"] = "../escape.xlsx"
    assert run_tool(client, office, agent, "excel.generate_sales", args)["status"] == "error"


def test_csv_conversion_and_folder_listing(client, tmp_path):
    office, agent, resource = setup_agent(client, tmp_path)
    (tmp_path / "entrada.csv").write_text("producto,total\nLápiz,1500\nCuaderno,3000\n", encoding="utf-8")
    result = run_tool(
        client,
        office,
        agent,
        "excel.from_csv",
        {
            "source_resource_id": resource["id"],
            "source_path": "entrada.csv",
            "resource_id": resource["id"],
            "path": "salida.xlsx",
        },
    )
    assert result["status"] == "completed", result
    book = load_workbook(tmp_path / "salida.xlsx")
    assert book["Datos"]["A2"].value == "Lápiz"
    book.close()
    result = run_tool(client, office, agent, "filesystem.list", {"resource_id": resource["id"], "limit": 1})
    assert result["status"] == "completed", result
    assert '"truncated": true' in result["result"]
    # Source read permission is separate from destination create permission.
    other = create(client, "resources", name="No autorizado", path=str(tmp_path))
    assert (
        run_tool(
            client,
            office,
            agent,
            "excel.from_csv",
            {
                "source_resource_id": other["id"],
                "source_path": "entrada.csv",
                "resource_id": resource["id"],
                "path": "denegado.xlsx",
            },
        )["status"]
        == "error"
    )
    assert not (tmp_path / "denegado.xlsx").exists()


def test_sql_calculations_export_directly_and_deny_writes(client, tmp_path):
    office, agent, resource = setup_agent(client, tmp_path)
    source = tmp_path / "ventas.db"
    with sqlite3.connect(source) as db:
        db.execute("CREATE TABLE sales (producto TEXT, total INTEGER)")
        db.executemany("INSERT INTO sales VALUES (?, ?)", [("Cuaderno", 1500), ("Cuaderno", 3000)])
    connection = create(client, "connections", name="SQL prueba", kind="sqlite", config={"path": str(source)})
    assigned = client.get(f"/api/agents/{agent['id']}/assignments").json()
    assigned["connections"] = [{"connection_id": connection["id"], "permissions": ["connect", "select"]}]
    client.put(f"/api/agents/{agent['id']}/assignments", json=assigned)
    metadata = run_tool(client, office, agent, "sql.describe", {"connection_id": connection["id"], "table": "sales"})
    assert metadata["status"] == "completed", metadata
    assert "producto" in metadata["result"]
    assert "INTEGER" in metadata["result"]
    args = {
        "connection_id": connection["id"],
        "query": "SELECT producto, SUM(total) AS ventas FROM sales GROUP BY producto",
        "resource_id": resource["id"],
        "path": "sql.xlsx",
    }
    result = run_tool(client, office, agent, "sql.export_excel", args)
    assert result["status"] == "completed", result
    book = load_workbook(tmp_path / "sql.xlsx")
    assert book["Datos"]["B2"].value == 4500
    book.close()
    args.update(query="DELETE FROM sales", path="bad.xlsx")
    assert run_tool(client, office, agent, "sql.export_excel", args)["status"] == "error"
    assert not (tmp_path / "bad.xlsx").exists()
    with sqlite3.connect(source) as db:
        assert db.execute("SELECT COUNT(*) FROM sales").fetchone()[0] == 2


def test_social_draft_never_claims_publication(client, tmp_path):
    office, agent, resource = setup_agent(client, tmp_path)
    result = run_tool(
        client,
        office,
        agent,
        "social.prepare_post",
        {
            "resource_id": resource["id"],
            "path": "linkedin.json",
            "platform": "linkedin",
            "text": "Resumen para revisar",
        },
    )
    assert result["status"] == "completed", result
    draft = json.loads((tmp_path / "linkedin.json").read_text(encoding="utf-8"))
    assert draft["published"] is False
    assert draft["requires_publish_approval"] is True
    assert draft["status"] == "draft"
